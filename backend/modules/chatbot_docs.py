"""
modules/chatbot_docs.py

Fuente de conocimiento del chatbot: toma de la sección "Base de Datos"
SOLO los archivos con modulo = "documentacion", extrae su texto, lo
divide en fragmentos y elige los más relevantes para cada pregunta.

El texto extraído se guarda en memoria (caché por id de archivo). Si el
servidor reinicia, se vuelve a extraer en la primera pregunta.
"""

import io
import logging
import re
import threading
import unicodedata
from collections import Counter

import database
import supabase_storage

logger = logging.getLogger("chatbot.docs")

MODULO_DOCUMENTACION = "documentacion"
EXTENSIONES_SOPORTADAS = (".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm")
MAX_BYTES_ARCHIVO = 15 * 1024 * 1024  # archivos más grandes se omiten

TAMANO_FRAGMENTO = 1200   # caracteres por fragmento
SOLAPE_FRAGMENTO = 150    # caracteres compartidos entre fragmentos vecinos

STOPWORDS = {
    "que", "como", "para", "por", "con", "los", "las", "una", "uno", "unos",
    "unas", "del", "the", "and", "este", "esta", "esto", "ese", "esa", "eso",
    "donde", "cual", "cuales", "cuando", "hay", "puedo", "puede", "debo",
    "tengo", "tiene", "son", "ser", "sus", "mas", "pero", "sin", "sobre",
    "entre", "hacer", "usar", "quiero", "necesito",
}

# Detecta encabezados al inicio de línea: "ARTÍCULO 1o.", "Artículo 48.", "PARÁGRAFO 2o."
RE_ENCABEZADO = re.compile(
    r"^[ \t>*#\"“]*(?P<tipo>art[ií]culo|par[aá]grafo)\s+(?P<num>\d+\s?[oº°]?)\.",
    re.IGNORECASE | re.MULTILINE,
)
RE_PAGINA_MD = re.compile(r"^#{1,3}\s*P[áa]gina\s+(\d+)\s*$", re.IGNORECASE | re.MULTILINE)

# archivo_id -> {"huella": (...), "fragmentos": [...]}
_cache: dict[int, dict] = {}
_lock = threading.Lock()


# ============================================================
# NORMALIZACIÓN
# ============================================================

def _normalizar(texto) -> str:
    if not texto:
        return ""
    texto = unicodedata.normalize("NFD", str(texto).lower())
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return texto.strip()


def _raiz(palabra: str) -> str:
    """Quita la 's' final de plurales simples (archivos -> archivo)."""
    return palabra[:-1] if len(palabra) > 3 and palabra.endswith("s") else palabra


def _palabras_clave(texto) -> list[str]:
    return [
        _raiz(p)
        for p in _normalizar(texto).split()
        if len(p) > 2 and p not in STOPWORDS
    ]


# ============================================================
# EXTRACCIÓN DE TEXTO  -> lista de (ubicacion | None, texto)
# ============================================================

    from pypdf import PdfReader

    lector = PdfReader(io.BytesIO(datos))
    secciones = []
    for numero, pagina in enumerate(lector.pages, start=1):
        texto = (pagina.extract_text() or "").strip()
        if texto:
            secciones.append((f"pág. {numero}", texto))
    return secciones
def _secciones_pdf(datos: bytes):
    import pymupdf

    secciones = []
    with pymupdf.open(stream=datos, filetype="pdf") as doc:
        for numero, pagina in enumerate(doc, start=1):
            texto = pagina.get_text().strip()
            if texto:
                secciones.append((f"pág. {numero}", texto))
    return secciones

def _secciones_docx(datos: bytes):
    from docx import Document

    doc = Document(io.BytesIO(datos))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    return [(None, "\n".join(partes))] if partes else []


def _secciones_xlsx(datos: bytes):
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    secciones = []
    try:
        for hoja in libro.worksheets:
            filas = []
            for fila in hoja.iter_rows(values_only=True):
                celdas = [str(c).strip() for c in fila if c is not None and str(c).strip()]
                if celdas:
                    filas.append(" | ".join(celdas))
            if filas:
                secciones.append((f"hoja {hoja.title}", "\n".join(filas)))
    finally:
        libro.close()
    return secciones


    texto = datos.decode("utf-8", errors="replace").strip()
    return [(None, texto)] if texto else []
def _secciones_texto(datos: bytes):
    texto = datos.decode("utf-8", errors="replace").strip()
    if not texto:
        return []
    partes = RE_PAGINA_MD.split(texto)
    if len(partes) > 1:
        secciones = []
        if partes[0].strip():
            secciones.append((None, partes[0].strip()))
        for num, cuerpo in zip(partes[1::2], partes[2::2]):
            if cuerpo.strip():
                secciones.append((f"pág. {num}", cuerpo.strip()))
        return secciones
    return [(None, texto)]

def _extraer_secciones(nombre: str, datos: bytes):
    ext = nombre.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return _secciones_pdf(datos)
    if ext == "docx":
        return _secciones_docx(datos)
    if ext in ("xlsx", "xlsm"):
        return _secciones_xlsx(datos)
    return _secciones_texto(datos)


# ============================================================
# FRAGMENTACIÓN
# ============================================================

    nombre = archivo["nombre_original"]
    datos = supabase_storage.descargar_bytes(archivo["ruta"])

    fragmentos = []
    for ubicacion, texto in _extraer_secciones(nombre, datos):
        for trozo in _dividir(texto):
            fragmentos.append({
                "archivo": nombre,
                "archivo_id": archivo["id"],
                "ubicacion": ubicacion,
                "texto": trozo,
                "_conteo": Counter(_palabras_clave(trozo)),
                "_nombre_claves": set(_palabras_clave(nombre)),
            })
    return fragmentos
def _limpiar(texto: str) -> str:
    return re.sub(r"[ \t]+", " ", texto).strip()


def _dividir(texto: str) -> list[tuple[int, int, str]]:
    """Recibe texto ya limpio. Devuelve (inicio, fin, trozo)."""
    fragmentos = []
    i = 0
    while i < len(texto):
        fin = min(i + TAMANO_FRAGMENTO, len(texto))
        if fin < len(texto):
            corte = texto.rfind("\n", i + TAMANO_FRAGMENTO // 2, fin)
            if corte == -1:
                corte = texto.rfind(". ", i + TAMANO_FRAGMENTO // 2, fin)
            if corte != -1:
                fin = corte + 1
        trozo = texto[i:fin].strip()
        if trozo:
            fragmentos.append((i, fin, trozo))
        if fin >= len(texto):
            break
        i = max(fin - SOLAPE_FRAGMENTO, i + 1)
    return fragmentos


def _marcas_estructura(texto: str) -> list[tuple[int, str, str]]:
    marcas = []
    for m in RE_ENCABEZADO.finditer(texto):
        tipo = "articulo" if m.group("tipo").lower().startswith("art") else "paragrafo"
        marcas.append((m.start(), tipo, re.sub(r"\s+", "", m.group("num"))))
    return marcas


def _aplicar_marca(estado: dict, tipo: str, num: str) -> None:
    if tipo == "articulo":
        estado["articulo"] = f"Artículo {num}"
        estado["paragrafo"] = None
    else:
        estado["paragrafo"] = f"Parágrafo {num}"


def _rotulo(estado: dict) -> str:
    return ", ".join(x for x in (estado["articulo"], estado["paragrafo"]) if x)


def _fragmentos_de_archivo(archivo: dict) -> list[dict]:
    nombre = archivo["nombre_original"]
    datos = supabase_storage.descargar_bytes(archivo["ruta"])

    estado = {"articulo": None, "paragrafo": None}  # se arrastra entre páginas
    fragmentos = []

    for pagina, texto in _extraer_secciones(nombre, datos):
        texto = _limpiar(texto)
        marcas = _marcas_estructura(texto)

        for inicio, fin, trozo in _dividir(texto):
            local = dict(estado)
            for pos, tipo, num in marcas:          # estado al inicio del fragmento
                if pos <= inicio:
                    _aplicar_marca(local, tipo, num)
            rotulos = [_rotulo(local)]
            for pos, tipo, num in marcas:          # encabezados dentro del fragmento
                if inicio < pos < fin:
                    _aplicar_marca(local, tipo, num)
                    rotulos.append(_rotulo(local))
            rotulos = [r for r in dict.fromkeys(rotulos) if r]

            partes = [p for p in (pagina, "; ".join(rotulos)) if p]
            fragmentos.append({
                "archivo": nombre,
                "archivo_id": archivo["id"],
                "ubicacion": " · ".join(partes) or None,
                "texto": trozo,
                "_conteo": Counter(_palabras_clave(trozo)),
                "_nombre_claves": set(_palabras_clave(nombre)),
            })

        for pos, tipo, num in marcas:              # estado al final de la página
            _aplicar_marca(estado, tipo, num)

    return fragmentos

# ============================================================
# API DEL MÓDULO
# ============================================================

def listar_documentos() -> list[dict]:
    """Archivos de la BD con modulo = 'documentacion' que el chatbot puede leer."""
    archivos = database.listar_archivos_por_modulo(MODULO_DOCUMENTACION)
    return [
        a for a in archivos
        if a["nombre_original"].lower().endswith(EXTENSIONES_SOPORTADAS)
        and not a["nombre_original"].startswith("~$")
    ]


def cargar_fragmentos() -> tuple[list[dict], list[dict], list[str]]:
    """
    Devuelve (fragmentos, documentos, fallidos).
    - Usa la caché; solo descarga/extrae archivos nuevos o cambiados.
    - Quita de la caché los archivos que ya no tienen módulo 'documentacion'
      o que fueron eliminados.
    """
    documentos = listar_documentos()
    vigentes = {a["id"] for a in documentos}

    with _lock:
        for archivo_id in list(_cache):
            if archivo_id not in vigentes:
                del _cache[archivo_id]

    fragmentos, fallidos = [], []

    for archivo in documentos:
        nombre = archivo["nombre_original"]
        huella = (archivo["id"], archivo.get("tamano_bytes"), archivo.get("fecha_subida"))

        if (archivo.get("tamano_bytes") or 0) > MAX_BYTES_ARCHIVO:
            logger.warning("Omitido por tamaño: %s", nombre)
            fallidos.append(nombre)
            continue

        with _lock:
            entrada = _cache.get(archivo["id"])

        if entrada is None or entrada["huella"] != huella:
            try:
                entrada = {"huella": huella, "fragmentos": _fragmentos_de_archivo(archivo)}
            except Exception:
                logger.exception("No se pudo procesar %s", nombre)
                fallidos.append(nombre)
                continue
            with _lock:
                _cache[archivo["id"]] = entrada
            logger.info("Indexado: %s (%d fragmentos)", nombre, len(entrada["fragmentos"]))

        fragmentos.extend(entrada["fragmentos"])

    return fragmentos, documentos, fallidos


def seleccionar_fragmentos(pregunta: str, fragmentos: list[dict], k: int = 5) -> list[dict]:
    """Los k fragmentos con más coincidencias de palabras con la pregunta."""
    claves = set(_palabras_clave(pregunta))
    if not claves:
        return []

    puntuados = []
    for f in fragmentos:
        puntaje = sum(min(f["_conteo"][c], 3) for c in claves)
        puntaje += 2 * len(claves & f["_nombre_claves"])
        if puntaje > 0:
            puntuados.append((puntaje, f))

    puntuados.sort(key=lambda par: par[0], reverse=True)
    return [f for _, f in puntuados[:k]]