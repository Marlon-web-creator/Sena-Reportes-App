"""
modules/correo_aprendices.py

Cruza el número de documento del consolidado contra varios archivos
.xls/.xlsx/.xlsm (columna B = documento, columna F = correo) y agrega una
columna "Correo Electrónico" a cada hoja del consolidado, INTEGRADA al
diseño existente: si la hoja tiene una tabla de Excel, la columna nueva
pasa a formar parte de la tabla (mismo estilo, filtros y bandas) y hereda
el formato de la columna vecina.
"""

from copy import copy
from pathlib import Path

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from modules.file_utils import cargar_workbook_compatible

TITULO_COLUMNA_CORREO = "Correo Electrónico"
ANCHO_COLUMNA_CORREO = 36


def _normalizar_documento(valor) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto


def construir_mapa_documento_correo(archivos_xls: list[Path], fila_inicio: int) -> tuple[dict, dict]:
    mapa: dict[str, str] = {}
    duplicados: dict[str, list[str]] = {}

    for archivo in archivos_xls:
        wb = cargar_workbook_compatible(archivo)
        for ws in wb.worksheets:
            for fila in range(fila_inicio, ws.max_row + 1):
                documento = _normalizar_documento(ws.cell(fila, 2).value)   # Columna B
                correo_valor = ws.cell(fila, 6).value                       # Columna F
                correo = str(correo_valor).strip() if correo_valor is not None else ""

                if not documento or not correo:
                    continue

                if documento not in mapa:
                    mapa[documento] = correo
                elif mapa[documento] != correo:
                    duplicados.setdefault(documento, [mapa[documento]])
                    if correo not in duplicados[documento]:
                        duplicados[documento].append(correo)

    return mapa, duplicados


def _ubicar_zona_de_datos(ws, col_documento: int, fila_inicio_consolidado: int):
    """
    Devuelve (tabla, fila_encabezado, primera_fila_datos, ultima_fila, ultima_col).
    - Si la hoja tiene una tabla de Excel que contiene la columna del
      documento, usa SUS límites (así el encabezado queda en la fila
      correcta, p. ej. la 5, sin depender de lo que se escriba en el formulario).
    - Si no hay tabla, usa fila_inicio_consolidado - 1 como encabezado.
    """
    for tabla in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(tabla.ref)
        if min_col <= col_documento <= max_col:
            return tabla, min_row, min_row + 1, max_row, max_col

    return None, fila_inicio_consolidado - 1, fila_inicio_consolidado, ws.max_row, ws.max_column


def _copiar_estilo(origen, destino):
    if origen.has_style:
        destino.font = copy(origen.font)
        destino.fill = copy(origen.fill)
        destino.border = copy(origen.border)
        destino.alignment = copy(origen.alignment)
        destino.number_format = origen.number_format
        destino.protection = copy(origen.protection)


def enriquecer_con_correos(
    consolidado: Path,
    archivos_xls: list[Path],
    salida_dir: Path,
    fila_inicio_consolidado: int = 2,
    col_documento_consolidado: int = 4,   # Columna D
    fila_inicio_xls: int = 2,
) -> dict:
    salida_dir.mkdir(parents=True, exist_ok=True)

    mapa_correos, duplicados = construir_mapa_documento_correo(archivos_xls, fila_inicio_xls)
    wb = cargar_workbook_compatible(consolidado)

    total_documentos = 0
    total_encontrados = 0
    documentos_sin_correo = []
    hojas_con_columna = 0

    for ws in wb.worksheets:
        tabla, fila_enc, fila_ini, fila_fin, ultima_col = _ubicar_zona_de_datos(
            ws, col_documento_consolidado, fila_inicio_consolidado
        )

        filas_con_documento = [
            f for f in range(fila_ini, fila_fin + 1)
            if _normalizar_documento(ws.cell(f, col_documento_consolidado).value)
        ]
        if not filas_con_documento:
            continue  # hoja sin registros (p. ej. "No hay juicios…"): no se le agrega la columna

        col_nueva = ultima_col + 1

        # Encabezado: hereda el formato del encabezado vecino (color, fuente, borde)
        celda_enc = ws.cell(fila_enc, col_nueva, TITULO_COLUMNA_CORREO)
        _copiar_estilo(ws.cell(fila_enc, ultima_col), celda_enc)
        if tabla is None and not celda_enc.has_style:
            celda_enc.font = Font(bold=True)   # sin tabla no hay estilo que lo ponga en negrita

        for fila in range(fila_ini, fila_fin + 1):
            celda = ws.cell(fila, col_nueva)
            _copiar_estilo(ws.cell(fila, ultima_col), celda)

            documento = _normalizar_documento(ws.cell(fila, col_documento_consolidado).value)
            if not documento:
                continue

            total_documentos += 1
            correo = mapa_correos.get(documento)
            if correo:
                celda.value = correo
                total_encontrados += 1
            else:
                documentos_sin_correo.append(documento)

        ws.column_dimensions[get_column_letter(col_nueva)].width = ANCHO_COLUMNA_CORREO

        # La columna pasa a ser parte de la tabla: mismo estilo, bandas y filtros
        if tabla is not None:
            min_col, min_row, _, max_row = range_boundaries(tabla.ref)
            tabla.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(col_nueva)}{max_row}"
            tabla.tableColumns = []      # openpyxl las regenera al guardar, con los textos del encabezado
            tabla.autoFilter = None

        hojas_con_columna += 1

    ruta_salida = salida_dir / (Path(consolidado).stem + "_ConCorreos.xlsx")
    wb.save(ruta_salida)

    return {
        "archivo_generado": str(ruta_salida),
        "hojas_con_columna": hojas_con_columna,
        "total_correos_indexados": len(mapa_correos),
        "total_documentos_consolidado": total_documentos,
        "total_encontrados": total_encontrados,
        "total_sin_correo": len(documentos_sin_correo),
        "documentos_sin_correo": documentos_sin_correo[:50],
        "documentos_con_correos_distintos": duplicados,
    }