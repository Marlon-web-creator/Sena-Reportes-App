"""
modules/file_utils.py

Utilidades compartidas para detectar y cargar archivos tipo Excel
(.xls, .xlsx, .xlsm) sin importar cuál de los tres formatos llegue.
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

EXTENSIONES_EXCEL = {".xls", ".xlsx", ".xlsm"}


def es_excel(nombre_archivo: str) -> bool:
    """True si el nombre de archivo tiene una extensión de Excel soportada."""
    return Path(nombre_archivo).suffix.lower() in EXTENSIONES_EXCEL


def cargar_workbook_compatible(path: Path) -> Workbook:
    """
    Carga cualquier archivo .xls / .xlsx / .xlsm como un Workbook de
    openpyxl, para poder aplicarle la misma lógica de marcado sin
    importar el formato de origen.

    - .xlsx / .xlsm: se abren directamente con openpyxl (conservan
      formato, colores, etc).
    - .xls (formato antiguo, no soportado por openpyxl): se leen con
      pandas/xlrd y se reconstruyen como un Workbook nuevo. Se
      conservan los valores de todas las celdas, pero no el formato
      original (no lo tenía relevancia para el marcado, que se aplica
      después de todas formas).
    """
    extension = path.suffix.lower()

    if extension in (".xlsx", ".xlsm"):
        return load_workbook(path, keep_vba=(extension == ".xlsm"))

    if extension == ".xls":
        import pandas as pd

        hojas = pd.read_excel(path, sheet_name=None, header=None, engine="xlrd")
        wb = Workbook()
        wb.remove(wb.active)
        for nombre_hoja, df in hojas.items():
            ws = wb.create_sheet(title=str(nombre_hoja)[:31])
            for i, fila in enumerate(df.itertuples(index=False), start=1):
                for j, valor in enumerate(fila, start=1):
                    if pd.isna(valor):
                        valor = None
                    ws.cell(row=i, column=j, value=valor)
        return wb

    raise ValueError(f"Formato no soportado: {extension}")