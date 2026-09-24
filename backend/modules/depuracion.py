"""
modules/depuracion.py

Marca en rojo (celdas de nombre/apellido y pestaña) a las personas que
superan `limite` juicios "POR EVALUAR" en cada hoja del archivo.

No depende de posiciones fijas: en cada hoja se busca la fila de
encabezados por el texto de las cabeceras ("Nombres", "Apellidos",
"Juicio de Evaluación" y, si existe, "Número de Documento") y desde ahí
se ubican las columnas. Cada hoja puede tener un orden de columnas
distinto.
"""

import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl.styles import PatternFill

from .file_utils import cargar_workbook_compatible, es_excel

RELLENO_ROJO = PatternFill(fill_type="solid", fgColor="FF0000")
COLOR_PESTANA_ROJO = "FF0000"

VALOR_POR_EVALUAR = "POR EVALUAR"
MAX_FILAS_BUSQUEDA_ENCABEZADO = 30

# Nombres posibles de cada columna (ya normalizados: sin tildes, minúsculas)
_ALIAS = {
    "nombres": {"nombres", "nombre", "nombre(s)"},
    "apellidos": {"apellidos", "apellido"},
    "documento": {"numero de documento", "numero documento", "nro documento",
                  "no documento", "documento", "identificacion", "cedula"},
}
_REQUERIDAS = ("nombres", "apellidos", "juicio")


def _normalizar_encabezado(valor) -> str:
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower().strip()
    return re.sub(r"\s+", " ", texto)


def _texto(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _documento(valor) -> str:
    texto = _texto(valor)
    return texto[:-2] if texto.endswith(".0") else texto


def _detectar_encabezado(ws, fila_forzada: int | None = None):
    """
    Devuelve (fila_encabezado, {'nombres': col, 'apellidos': col, 'juicio': col,
    'documento': col (opcional)}) o None si la hoja no tiene la estructura esperada.
    """
    filas = [fila_forzada] if fila_forzada else range(1, min(ws.max_row, MAX_FILAS_BUSQUEDA_ENCABEZADO) + 1)

    for fila in filas:
        cols: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            texto = _normalizar_encabezado(ws.cell(fila, col).value)
            if not texto:
                continue
            for clave, alias in _ALIAS.items():
                if texto in alias and clave not in cols:
                    cols[clave] = col
            # "Juicio de Evaluación", "Juicio evaluativo"... (empieza por "juicio";
            # así no se confunde con "Funcionario que registró el juicio evaluativo")
            if texto.startswith("juicio") and "juicio" not in cols:
                cols["juicio"] = col

        if all(c in cols for c in _REQUERIDAS):
            return fila, cols

    return None


def depurar(archivo: Path, limite: int, salida_dir: Path, fila_encabezado: int | None = None) -> dict:
    """
    Marca en rojo (celda y pestaña) a las personas que superan `limite`
    juicios "POR EVALUAR" en cada hoja del archivo. Guarda el resultado
    en salida_dir y devuelve estadísticas + ruta del archivo generado.

    fila_encabezado es opcional: si no se indica se detecta en cada hoja.
    Las hojas sin las columnas necesarias se omiten y se listan en
    "hojas_omitidas".
    """
    archivo = Path(archivo)

    if not es_excel(archivo.name):
        raise ValueError(f"Formato no soportado: {archivo.suffix}")

    salida_dir.mkdir(parents=True, exist_ok=True)

    wb = cargar_workbook_compatible(archivo)

    total_personas_marcadas = 0
    total_hojas_marcadas = 0
    detalle_hojas = []
    hojas_omitidas = []

    for ws in wb.worksheets:
        ubicacion = _detectar_encabezado(ws, fila_encabezado)
        if ubicacion is None:
            hojas_omitidas.append(ws.title)
            continue

        fila_enc, cols = ubicacion
        col_nombre, col_apellido, col_juicio = cols["nombres"], cols["apellidos"], cols["juicio"]
        col_documento = cols.get("documento")
        filas_datos = range(fila_enc + 1, ws.max_row + 1)

        def clave_persona(fila: int) -> tuple | None:
            """
            Identifica a la persona por documento si la hoja lo tiene (evita
            juntar a dos personas con el mismo nombre); si no, por nombre y apellido.
            """
            if col_documento:
                doc = _documento(ws.cell(fila, col_documento).value)
                if doc:
                    return ("doc", doc)
            nombre = _texto(ws.cell(fila, col_nombre).value).upper()
            apellido = _texto(ws.cell(fila, col_apellido).value).upper()
            return ("nom", nombre, apellido) if nombre and apellido else None

        # Primer recorrido: contar "POR EVALUAR" por persona
        por_evaluar = defaultdict(int)
        for fila in filas_datos:
            if _texto(ws.cell(fila, col_juicio).value).upper() != VALOR_POR_EVALUAR:
                continue
            persona = clave_persona(fila)
            if persona:
                por_evaluar[persona] += 1

        personas_sobre_limite = {p for p, n in por_evaluar.items() if n > limite}

        hoja_marcada = bool(personas_sobre_limite)
        if hoja_marcada:
            ws.sheet_properties.tabColor = COLOR_PESTANA_ROJO
            total_hojas_marcadas += 1

        # Segundo recorrido: pintar celdas de nombre y apellido
        marcadas_en_hoja = set()
        for fila in filas_datos:
            persona = clave_persona(fila)
            if persona in personas_sobre_limite:
                ws.cell(fila, col_nombre).fill = RELLENO_ROJO
                ws.cell(fila, col_apellido).fill = RELLENO_ROJO
                marcadas_en_hoja.add(persona)

        total_personas_marcadas += len(marcadas_en_hoja)

        detalle_hojas.append({
            "hoja": ws.title,
            "marcada": hoja_marcada,
            "personas_marcadas": len(marcadas_en_hoja),
        })

    if len(hojas_omitidas) == len(wb.worksheets):
        raise ValueError(
            "No se encontraron las columnas 'Nombres', 'Apellidos' y 'Juicio de Evaluación' "
            "en ninguna hoja del archivo."
        )

    ruta_salida = salida_dir / (archivo.stem + "_Procesado.xlsx")
    wb.save(ruta_salida)

    return {
        "limite": limite,
        "total_hojas_marcadas": total_hojas_marcadas,
        "total_personas_marcadas": total_personas_marcadas,
        "detalle_hojas": detalle_hojas,
        "hojas_omitidas": hojas_omitidas,
        "archivo_generado": str(ruta_salida),
    }