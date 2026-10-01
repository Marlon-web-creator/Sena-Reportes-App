"""
modules/horas_instructores.py

Compara las horas de los instructores de un "Reporte Ejecución Horas
Instructor" (el archivo que se sube en cada ejecución) contra los Excels de
horas programadas que están en la sección "Base de Datos" con módulo
"Horas Instructores" (ej. HORAS_INSTRUCTORES_CONTRATISTA.xlsx, hoja "Export").

Hay dos tipos de Excel de horas programadas en la BD:

    CONTRATISTA  trae la columna "TIPOS DE HORAS AGRUPADAS" (HORAS ASOCIADAS A
                 GRUPOS / INASISTENCIA / OTRAS HORAS).
    PLANTA       (HORAS_INSTRUCTORES_PLANTA.xlsx) NO trae esa columna: cada fila
                 es la programación de un grupo, así que todas sus horas se
                 cuentan como "HORAS ASOCIADAS A GRUPOS".

Qué se compara (por número de documento, para UN mes elegido):

    Reporte  : "Horas Formación Titulada Etapa Lectiva" (+ "Etapa Productiva"
               si viene en el reporte)
    BD       : tipo "HORAS ASOCIADAS A GRUPOS" del mes elegido

La BD solo contiene programas TITULADA, por eso NO se usa la formación
complementaria del reporte. Las horas de INASISTENCIA y OTRAS HORAS de la BD,
y la formación complementaria, las horas adicionales y el total del reporte,
NO entran en la diferencia: se muestran como columnas informativas.

El reporte no trae el mes, por eso el mes se recibe como parámetro.

Además de la diferencia, el Excel de salida indica para cada instructor la
ACCIÓN a seguir y las HORAS POR SUBIR (lo que le falta al reporte para llegar a
lo programado en la BD).

Estados posibles por instructor:
    COINCIDE        horas iguales en ambos
    REPORTE > BD    el reporte tiene más horas que la BD
    REPORTE < BD    el reporte tiene menos horas que la BD
    SOLO EN REPORTE el documento no aparece en la BD para ese mes
    SOLO EN BD      el documento no aparece en el reporte
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ProgressCallback = Optional[Callable[[str, int, int], None]]

MESES = {
    1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO",
    7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE",
    12: "DICIEMBRE",
}

EXTENSIONES_BD = (".xlsx", ".xlsm")

# Estados (texto exacto que se escribe en el Excel)
COINCIDE = "COINCIDE"
REPORTE_MAYOR = "REPORTE > BD"
REPORTE_MENOR = "REPORTE < BD"
SOLO_REPORTE = "SOLO EN REPORTE"
SOLO_BD = "SOLO EN BD"

_ORDEN_ESTADOS = {
    REPORTE_MAYOR: 0, REPORTE_MENOR: 1, SOLO_BD: 2, SOLO_REPORTE: 3, COINCIDE: 4,
}

_COLOR_ESTADO = {
    COINCIDE: "E2EFDA",
    REPORTE_MAYOR: "FFF2CC",
    REPORTE_MENOR: "FCE4D6",
    SOLO_REPORTE: "DDEBF7",
    SOLO_BD: "F8CBAD",
}

# Columnas requeridas (ya normalizadas, ver _normalizar)
_REQ_REPORTE = ("NUMERO IDENTIFICACION", "HORAS FORMACION TITULADA ETAPA LECTIVA")
_REQ_BD = ("DOCUMENTO", "MES PROGRAMADO", "TOTAL HORAS")
_COL_TIPO_BD = "TIPOS DE HORAS AGRUPADAS"  # opcional: los Excels de PLANTA no la traen
_TIPO_POR_DEFECTO = "HORAS ASOCIADAS A GRUPOS"

# Acciones (texto exacto que se escribe en el Excel)
ACCION_OK = "OK"
ACCION_SUBIR = "SUBIR HORAS"
ACCION_SUBIR_TODAS = "SUBIR HORAS (no aparece en el reporte)"
ACCION_REVISAR_EXCESO = "REVISAR: el reporte tiene más horas que la BD"
ACCION_REVISAR_SOLO_REPORTE = "REVISAR: no tiene horas programadas en la BD"


# ============================================================
# Utilidades
# ============================================================

def _notificar(cb: ProgressCallback, mensaje: str, actual: int, total: int) -> None:
    if cb:
        cb(mensaje, actual, total)


def _normalizar(texto) -> str:
    """Mayúsculas, sin tildes y con espacios colapsados (para comparar encabezados)."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    s = unicodedata.normalize("NFKD", str(texto))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


def _normalizar_documento(valor) -> Optional[str]:
    """
    Deja el documento como texto de solo dígitos. Soporta 1012336732,
    1012336732.0 (float de Excel), '1.012.336.732' y '1012336732 '.
    """
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, int):
        return str(valor)
    s = str(valor).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    s = re.sub(r"\D", "", s)
    return s or None


def _a_horas(valor) -> float:
    """Convierte a número; vacío o texto no numérico cuenta como 0."""
    n = pd.to_numeric(valor, errors="coerce")
    return 0.0 if pd.isna(n) else float(n)


def _numero_mes(valor) -> Optional[int]:
    """'9. SEPTIEMBRE' -> 9 ; 'SEPTIEMBRE' -> 9 ; otro -> None."""
    s = _normalizar(valor)
    if not s:
        return None
    m = re.match(r"(\d{1,2})\b", s)
    if m:
        n = int(m.group(1))
        return n if 1 <= n <= 12 else None
    for n, nombre in MESES.items():
        if nombre in s:
            return n
    return None


def _num(x) -> float | int:
    """float entero -> int (para que el JSON y el Excel no muestren '176.0')."""
    x = float(x)
    return int(x) if x.is_integer() else round(x, 2)


def _origen_archivo(nombre: str) -> str:
    """PLANTA / CONTRATISTA según el nombre del Excel; vacío si no se reconoce."""
    n = _normalizar(nombre)
    if "PLANTA" in n:
        return "PLANTA"
    if "CONTRATISTA" in n or "CONTRATO" in n:
        return "CONTRATISTA"
    return ""


def _detectar_encabezado(crudo: pd.DataFrame, requeridas: tuple, max_filas: int = 40):
    """
    Busca la fila que contiene TODAS las columnas requeridas. Devuelve
    (indice_fila, {nombre_normalizado: indice_columna}) o (None, {}).
    """
    for i in range(min(max_filas, len(crudo))):
        mapa = {}
        for j, v in enumerate(crudo.iloc[i].tolist()):
            n = _normalizar(v)
            if n and n not in mapa:
                mapa[n] = j
        if all(r in mapa for r in requeridas):
            return i, mapa
    return None, {}


# ============================================================
# Lectura del reporte subido
# ============================================================

def leer_reporte(ruta: Path) -> tuple[pd.DataFrame, list[str]]:
    """
    Devuelve (df, documentos_repetidos). df tiene una fila por documento:
    doc, nombre_reporte, vinculacion, titulada_reporte (lectiva + productiva),
    complementaria_reporte, adicionales_reporte, total_reporte. Si un documento viene repetido en el reporte se suman sus
    horas y se avisa.
    """
    crudo = pd.read_excel(ruta, sheet_name=0, header=None)
    fila, mapa = _detectar_encabezado(crudo, _REQ_REPORTE)
    if fila is None:
        raise ValueError(
            "El reporte no tiene las columnas esperadas "
            "('NUMERO IDENTIFICACION' y 'Horas Formación Titulada Etapa Lectiva'). "
            "¿Es el 'Reporte Ejecución Horas Instructor'?"
        )

    datos = crudo.iloc[fila + 1:]

    def col(nombre):
        j = mapa.get(nombre)
        return datos.iloc[:, j] if j is not None else pd.Series([None] * len(datos), index=datos.index)

    df = pd.DataFrame({
        "doc": col("NUMERO IDENTIFICACION").map(_normalizar_documento),
        "apellido": col("APELLIDO"),
        "nombres": col("NOMBRES"),
        "vinculacion": col("TIPO VINCULACION"),
        "lectiva": col("HORAS FORMACION TITULADA ETAPA LECTIVA").map(_a_horas),
        "productiva": col("HORAS FORMACION TITULADA ETAPA PRODUCTIVA").map(_a_horas),
        "complementaria_reporte": col("HORAS FORMACION COMPLEMENTARIA").map(_a_horas),
        "adicionales_reporte": col("TOTAL HORAS ADICIONALES").map(_a_horas),
        "total_reporte": col("TOTAL HORAS INSTRUCTOR").map(_a_horas),
    })
    df = df[df["doc"].notna()].copy()
    df["titulada_reporte"] = df["lectiva"] + df["productiva"]

    def _nombre(r):
        partes = [str(p).strip() for p in (r["apellido"], r["nombres"]) if pd.notna(p) and str(p).strip()]
        return " ".join(partes)

    df["nombre_reporte"] = df.apply(_nombre, axis=1)

    repetidos = sorted(df.loc[df["doc"].duplicated(keep=False), "doc"].unique().tolist())
    if repetidos:
        df = df.groupby("doc", as_index=False).agg(
            nombre_reporte=("nombre_reporte", "first"),
            vinculacion=("vinculacion", "first"),
            titulada_reporte=("titulada_reporte", "sum"),
            complementaria_reporte=("complementaria_reporte", "sum"),
            adicionales_reporte=("adicionales_reporte", "sum"),
            total_reporte=("total_reporte", "sum"),
        )

    cols = ["doc", "nombre_reporte", "vinculacion", "titulada_reporte",
            "complementaria_reporte", "adicionales_reporte", "total_reporte"]
    return df[cols].reset_index(drop=True), repetidos


# ============================================================
# Lectura de los Excels de la Base de Datos
# ============================================================

def _leer_bd_archivo(ruta: Path, mes: int, omitidos: list[dict]):
    """
    Lee un Excel de la BD. Devuelve (df_agregado_por_doc | None,
    meses_encontrados:set[int], tipos_no_reconocidos:set[str]).
    Prueba todas las hojas y usa las que tengan las columnas esperadas.
    """
    hojas = pd.read_excel(ruta, sheet_name=None, header=None)
    partes = []
    meses_encontrados: set[int] = set()
    tipos_raros: set[str] = set()
    alguna_valida = False

    for nombre_hoja, crudo in hojas.items():
        fila, mapa = _detectar_encabezado(crudo, _REQ_BD)
        if fila is None:
            omitidos.append({
                "archivo": ruta.name, "hoja": str(nombre_hoja),
                "motivo": "no tiene las columnas DOCUMENTO / MES PROGRAMADO / TOTAL HORAS",
            })
            continue
        alguna_valida = True
        datos = crudo.iloc[fila + 1:]

        def col(nombre):
            j = mapa.get(nombre)
            return datos.iloc[:, j] if j is not None else pd.Series([None] * len(datos), index=datos.index)

        df = pd.DataFrame({
            "doc": col("DOCUMENTO").map(_normalizar_documento),
            "nombre_bd": col("NOMBRE INSTRUCTOR"),
            # Excels de PLANTA: sin columna de tipo -> todo son horas asociadas a grupos
            "tipo": (
                col(_COL_TIPO_BD).map(_normalizar)
                if _COL_TIPO_BD in mapa
                else pd.Series(_TIPO_POR_DEFECTO, index=datos.index)
            ),
            "mes": col("MES PROGRAMADO").map(_numero_mes),
            "horas": col("TOTAL HORAS").map(_a_horas),
        })
        # Descarta filas de "Total", pie de filtros y filas vacías (sin documento/tipo/mes);
        # la fila "Total" y el pie de filtros no traen documento, por eso salen aquí.
        df = df[df["doc"].notna() & df["tipo"].ne("") & df["mes"].notna()]
        meses_encontrados |= set(int(m) for m in df["mes"].unique())
        df = df[df["mes"] == mes].copy()
        if df.empty:
            continue

        def clasificar(t: str) -> Optional[str]:
            if "ASOCIADAS" in t:
                return "grupos_bd"
            if "INASISTENCIA" in t:
                return "inasistencia_bd"
            if t.startswith("OTRAS"):
                return "otras_bd"
            return None

        df["campo"] = df["tipo"].map(clasificar)
        tipos_raros |= set(df.loc[df["campo"].isna(), "tipo"].unique())
        df = df[df["campo"].notna()]
        if df.empty:
            continue

        horas = df.pivot_table(index="doc", columns="campo", values="horas",
                               aggfunc="sum", fill_value=0.0)
        for c in ("grupos_bd", "inasistencia_bd", "otras_bd"):
            if c not in horas.columns:
                horas[c] = 0.0
        nombres = df.dropna(subset=["nombre_bd"]).groupby("doc")["nombre_bd"].first()
        horas["nombre_bd"] = nombres
        partes.append(horas[["nombre_bd", "grupos_bd", "inasistencia_bd", "otras_bd"]].reset_index())

    if not alguna_valida:
        return None, meses_encontrados, tipos_raros
    if not partes:
        return pd.DataFrame(columns=["doc", "nombre_bd", "grupos_bd", "inasistencia_bd", "otras_bd"]), meses_encontrados, tipos_raros
    return pd.concat(partes, ignore_index=True), meses_encontrados, tipos_raros


def cargar_bd(carpeta_bd: Path, mes: int, progress_callback: ProgressCallback = None):
    """
    Lee todos los Excels de `carpeta_bd` (recursivo). Devuelve
    (df_bd_por_doc, info) con info = archivos_usados, omitidos,
    documentos_en_varios_archivos, tipos_no_reconocidos, meses_disponibles.
    """
    archivos = sorted(
        p for p in Path(carpeta_bd).rglob("*")
        if p.is_file() and p.suffix.lower() in EXTENSIONES_BD and not p.name.startswith("~$")
    )
    omitidos: list[dict] = []
    usados: list[str] = []
    todos_meses: set[int] = set()
    tipos_raros: set[str] = set()
    por_archivo = []

    total = len(archivos)
    for n, ruta in enumerate(archivos, start=1):
        _notificar(progress_callback, f"Leyendo Excels de la BD: {n}/{total}…", n, total)
        try:
            df, meses, raros = _leer_bd_archivo(ruta, mes, omitidos)
        except Exception as e:  # archivo ilegible: se omite y se avisa
            omitidos.append({"archivo": ruta.name, "hoja": "", "motivo": f"no se pudo leer ({type(e).__name__}: {e})"})
            continue
        todos_meses |= meses
        tipos_raros |= raros
        if df is None:
            continue
        usados.append(ruta.name)
        if not df.empty:
            df = df.copy()
            df["archivo"] = ruta.name
            df["origen"] = _origen_archivo(ruta.name)
            por_archivo.append(df)

    info = {
        "archivos_usados": usados,
        "omitidos": omitidos,
        "tipos_no_reconocidos": sorted(tipos_raros),
        "meses_disponibles": sorted(todos_meses),
        "documentos_en_varios_archivos": [],
    }

    vacio = pd.DataFrame(columns=["doc", "nombre_bd", "grupos_bd", "inasistencia_bd", "otras_bd", "archivos_bd", "origen_bd"])
    if not por_archivo:
        return vacio, info

    junto = pd.concat(por_archivo, ignore_index=True)
    varios = (
        junto.groupby("doc")["archivo"].nunique().loc[lambda s: s > 1].index.tolist()
    )
    info["documentos_en_varios_archivos"] = sorted(varios)

    agregado = junto.groupby("doc", as_index=False).agg(
        nombre_bd=("nombre_bd", "first"),
        grupos_bd=("grupos_bd", "sum"),
        inasistencia_bd=("inasistencia_bd", "sum"),
        otras_bd=("otras_bd", "sum"),
        archivos_bd=("archivo", lambda s: ", ".join(sorted(set(s)))),
        origen_bd=("origen", lambda s: ", ".join(sorted({x for x in s if x}))),
    )
    return agregado, info


# ============================================================
# Comparación
# ============================================================

def comparar(df_reporte: pd.DataFrame, df_bd: pd.DataFrame) -> pd.DataFrame:
    m = df_reporte.merge(df_bd, on="doc", how="outer", indicator=True)

    for c in ("titulada_reporte", "complementaria_reporte", "adicionales_reporte",
              "total_reporte", "grupos_bd", "inasistencia_bd", "otras_bd"):
        m[c] = m[c].astype(float)
    if "origen_bd" not in m.columns:
        m["origen_bd"] = None

    def estado(r):
        if r["_merge"] == "left_only":
            return SOLO_REPORTE
        if r["_merge"] == "right_only":
            return SOLO_BD
        d = round(r["titulada_reporte"] - r["grupos_bd"], 2)
        if d == 0:
            return COINCIDE
        return REPORTE_MAYOR if d > 0 else REPORTE_MENOR

    m["estado"] = m.apply(estado, axis=1)
    ambos = m["_merge"] == "both"
    m["diferencia"] = None
    m.loc[ambos, "diferencia"] = (m.loc[ambos, "titulada_reporte"] - m.loc[ambos, "grupos_bd"]).round(2)

    # Qué hacer con cada instructor y cuántas horas le faltan al reporte
    _acciones = {
        COINCIDE: ACCION_OK,
        REPORTE_MENOR: ACCION_SUBIR,
        REPORTE_MAYOR: ACCION_REVISAR_EXCESO,
        SOLO_REPORTE: ACCION_REVISAR_SOLO_REPORTE,
        SOLO_BD: ACCION_SUBIR_TODAS,
    }
    m["accion"] = m["estado"].map(_acciones)

    def _por_subir(r):
        if r["estado"] == REPORTE_MENOR:
            return round(-float(r["diferencia"]), 2)
        if r["estado"] == SOLO_BD:
            return round(float(r["grupos_bd"]), 2)
        return 0.0

    m["horas_por_subir"] = m.apply(_por_subir, axis=1)

    m["_orden"] = m["estado"].map(_ORDEN_ESTADOS)
    m["_abs"] = m["diferencia"].map(lambda x: abs(float(x)) if x is not None and pd.notna(x) else 0.0)
    m["_nombre"] = m["nombre_reporte"].fillna(m["nombre_bd"]).fillna("")
    m = m.sort_values(["_orden", "_abs", "_nombre"], ascending=[True, False, True]).reset_index(drop=True)
    return m.drop(columns=["_merge", "_orden", "_abs", "_nombre"])


# ============================================================
# Excel de salida
# ============================================================

_BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
_ENCABEZADO_FILL = PatternFill("solid", fgColor="1F4E78")
_ENCABEZADO_FONT = Font(bold=True, color="FFFFFF")


def _vacio(v):
    return None if v is None or (isinstance(v, float) and pd.isna(v)) else v


def _escribir_excel(df: pd.DataFrame, resumen_filas: list[tuple], ruta: Path) -> None:
    wb = Workbook()

    # ---- Resumen
    ws = wb.active
    ws.title = "Resumen"
    for fila in resumen_filas:
        ws.append(list(fila))
    for celda in ws["A"]:
        if celda.value and ws.cell(row=celda.row, column=2).value is None:
            celda.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 60

    # ---- Comparativo
    ws2 = wb.create_sheet("Comparativo")
    encabezados = [
        "DOCUMENTO", "NOMBRE (REPORTE)", "NOMBRE (BD)", "TIPO VINCULACIÓN",
        "ORIGEN (BD)",
        "HORAS FORMACIÓN TITULADA (REPORTE)", "HORAS ASOCIADAS A GRUPOS (BD)",
        "DIFERENCIA (REPORTE - BD)", "ESTADO", "HORAS POR SUBIR", "ACCIÓN",
        "HORAS INASISTENCIA (BD)", "OTRAS HORAS (BD)",
        "HORAS FORMACIÓN COMPLEMENTARIA (REPORTE)", "HORAS ADICIONALES (REPORTE)",
        "TOTAL HORAS INSTRUCTOR (REPORTE)", "ARCHIVO(S) BD",
    ]
    ws2.append(encabezados)
    for c in ws2[1]:
        c.font = _ENCABEZADO_FONT
        c.fill = _ENCABEZADO_FILL
        c.border = _BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.row_dimensions[1].height = 42

    def num(v):
        v = _vacio(v)
        return None if v is None else _num(v)

    for _, r in df.iterrows():
        solo_bd = r["estado"] == SOLO_BD
        solo_rep = r["estado"] == SOLO_REPORTE
        fila = [
            r["doc"],
            _vacio(r["nombre_reporte"]),
            _vacio(r["nombre_bd"]),
            _vacio(r["vinculacion"]),
            _vacio(r["origen_bd"]) or None,
            None if solo_bd else num(r["titulada_reporte"]),
            None if solo_rep else num(r["grupos_bd"]),
            num(r["diferencia"]),
            r["estado"],
            num(r["horas_por_subir"]),
            r["accion"],
            None if solo_rep else num(r["inasistencia_bd"]),
            None if solo_rep else num(r["otras_bd"]),
            None if solo_bd else num(r["complementaria_reporte"]),
            None if solo_bd else num(r["adicionales_reporte"]),
            None if solo_bd else num(r["total_reporte"]),
            _vacio(r["archivos_bd"]),
        ]
        ws2.append(fila)
        color = PatternFill("solid", fgColor=_COLOR_ESTADO[r["estado"]])
        for c in ws2[ws2.max_row]:
            c.border = _BORDE
        for col_color in (8, 9, 10, 11):  # diferencia, estado, horas por subir, acción
            ws2.cell(row=ws2.max_row, column=col_color).fill = color
        ws2.cell(row=ws2.max_row, column=1).number_format = "@"

    anchos = [15, 34, 34, 26, 15, 18, 18, 16, 18, 14, 40, 14, 12, 20, 16, 18, 38]
    for i, a in enumerate(anchos, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = a
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{ws2.max_row}"

    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


# ============================================================
# Punto de entrada
# ============================================================

def procesar(
    archivo_reporte: Path,
    carpeta_bd: Path,
    salida_dir: Path,
    mes: int,
    progress_callback: ProgressCallback = None,
) -> dict:
    """
    Compara el reporte subido contra los Excels de `carpeta_bd` para el mes
    `mes` (1-12) y escribe `Comparativo_Horas_<MES>.xlsx` en `salida_dir`.
    Lanza ValueError con un mensaje legible si algo no se puede procesar.
    """
    if mes not in MESES:
        raise ValueError("El mes debe estar entre 1 y 12.")
    nombre_mes = MESES[mes]

    _notificar(progress_callback, "Leyendo el reporte…", 0, 1)
    df_reporte, repetidos_reporte = leer_reporte(Path(archivo_reporte))

    df_bd, info = cargar_bd(Path(carpeta_bd), mes, progress_callback)

    if not info["archivos_usados"]:
        raise ValueError(
            "Ninguno de los Excels de la Base de Datos tiene la estructura esperada "
            "(DOCUMENTO, MES PROGRAMADO, TOTAL HORAS)."
        )
    if df_bd.empty:
        disp = ", ".join(f"{n} ({MESES[n].capitalize()})" for n in info["meses_disponibles"]) or "ninguno"
        raise ValueError(
            f"Los Excels de la Base de Datos no tienen horas para {nombre_mes.capitalize()}. "
            f"Meses disponibles: {disp}."
        )

    _notificar(progress_callback, "Comparando instructores…", 0, 1)
    df = comparar(df_reporte, df_bd)

    conteo = df["estado"].value_counts().to_dict()
    solo_rep = df[df["estado"] == SOLO_REPORTE]
    por_vinc = (
        solo_rep["vinculacion"].fillna("(sin dato)").value_counts().to_dict()
        if not solo_rep.empty else {}
    )

    por_subir = df[df["horas_por_subir"] > 0]
    n_por_subir = int(len(por_subir))
    horas_por_subir = _num(por_subir["horas_por_subir"].sum())

    horas_rep = _num(df_reporte["titulada_reporte"].sum())
    horas_bd = _num(df_bd["grupos_bd"].sum())

    resultado = {
        "archivo_generado": "",
        "mes": mes,
        "mes_nombre": nombre_mes,
        "total_instructores_reporte": int(len(df_reporte)),
        "total_instructores_bd": int(len(df_bd)),
        "total_coinciden": int(conteo.get(COINCIDE, 0)),
        "total_reporte_mayor": int(conteo.get(REPORTE_MAYOR, 0)),
        "total_reporte_menor": int(conteo.get(REPORTE_MENOR, 0)),
        "total_solo_reporte": int(conteo.get(SOLO_REPORTE, 0)),
        "total_solo_bd": int(conteo.get(SOLO_BD, 0)),
        "horas_titulada_reporte": horas_rep,
        "horas_grupos_bd": horas_bd,
        "total_instructores_por_subir": n_por_subir,
        "total_horas_por_subir": horas_por_subir,
        "solo_reporte_por_vinculacion": {str(k): int(v) for k, v in por_vinc.items()},
        "archivos_bd_usados": info["archivos_usados"],
        "archivos_bd_omitidos": info["omitidos"],
        "documentos_en_varios_archivos": info["documentos_en_varios_archivos"],
        "documentos_repetidos_en_reporte": repetidos_reporte,
        "tipos_no_reconocidos": info["tipos_no_reconocidos"],
    }

    resumen = [
        ("COMPARATIVO DE HORAS DE INSTRUCTORES",),
        ("Mes comparado", nombre_mes),
        ("Reporte", Path(archivo_reporte).name),
        ("Excels de la BD usados", ", ".join(info["archivos_usados"])),
        ("",),
        ("Cómo se compara",),
        ("Reporte", "Horas Formación Titulada (Etapa Lectiva + Etapa Productiva)"),
        ("BD", "HORAS ASOCIADAS A GRUPOS (del mes)"),
        ("",),
        ("Resultado",),
        ("Instructores en el reporte", resultado["total_instructores_reporte"]),
        ("Instructores en la BD (mes)", resultado["total_instructores_bd"]),
        (COINCIDE, resultado["total_coinciden"]),
        (REPORTE_MAYOR, resultado["total_reporte_mayor"]),
        (REPORTE_MENOR, resultado["total_reporte_menor"]),
        (SOLO_REPORTE, resultado["total_solo_reporte"]),
        (SOLO_BD, resultado["total_solo_bd"]),
        ("",),
        ("Instructores a los que hay que subir horas", n_por_subir),
        ("Total de horas por subir", horas_por_subir),
        ("",),
        ("Total horas formación titulada (todo el reporte)", horas_rep),
        ("Total horas asociadas a grupos (toda la BD del mes)", horas_bd),
    ]
    if por_vinc:
        resumen += [("",), ("SOLO EN REPORTE por tipo de vinculación",)]
        resumen += [(k, int(v)) for k, v in por_vinc.items()]
    avisos = []
    if info["omitidos"]:
        avisos += [(f"Omitido: {o['archivo']}" + (f" [{o['hoja']}]" if o["hoja"] else ""), o["motivo"]) for o in info["omitidos"]]
    if info["documentos_en_varios_archivos"]:
        avisos.append(("Documentos presentes en varios Excels de la BD (se sumaron)", ", ".join(info["documentos_en_varios_archivos"])))
    if repetidos_reporte:
        avisos.append(("Documentos repetidos en el reporte (se sumaron)", ", ".join(repetidos_reporte)))
    if info["tipos_no_reconocidos"]:
        avisos.append(("Tipos de hora de la BD no reconocidos (ignorados)", ", ".join(info["tipos_no_reconocidos"])))
    if avisos:
        resumen += [("",), ("Avisos",)] + avisos

    ruta_salida = Path(salida_dir) / f"Comparativo_Horas_{nombre_mes}.xlsx"
    _notificar(progress_callback, "Generando el Excel…", 0, 1)
    _escribir_excel(df, resumen, ruta_salida)

    resultado["archivo_generado"] = str(ruta_salida)
    _notificar(progress_callback, "Listo.", 1, 1)
    return resultado