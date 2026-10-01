"""
modules/juicios_practica.py

Verifica, ficha por ficha y aprendiz por aprendiz, si existen juicios
evaluativos APROBADOS de la fase lectiva con fecha POSTERIOR a la fecha de
aprobación del juicio de la etapa práctica
("2 - RESULTADOS DE APRENDIZAJE ETAPA PRACTICA").

Regla de negocio: el resultado de etapa práctica se registra al final del
proceso, así que no debería haber juicios lectivos aprobados después de él.
Todo juicio aprobado con fecha mayor a la de la práctica es una
inconsistencia a revisar.

Entrada: uno o varios "Reporte de Juicios de Evaluación" (.xls/.xlsx/.xlsm),
uno por ficha. Igual que en correo_aprendices.py, nada se recibe por
posición fija; todo se detecta del contenido:

  - La fila de encabezado se localiza puntuando cada fila según cuántos
    campos conocidos contiene (Número de Documento, Competencia, Resultado
    de Aprendizaje, Juicio de Evaluación, Fecha y Hora del Juicio
    Evaluativo, ...). Los datos empiezan justo debajo.
  - Cada columna se elige por el texto de su encabezado.
  - El número de ficha se toma de la cabecera del reporte
    ("Ficha de Caracterización:"); si no está, se usa el nombre del archivo.
  - La competencia de etapa práctica se identifica por el texto "ETAPA
    PRACTICA" en la columna Competencia.
"""

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from modules.file_utils import cargar_workbook_compatible

MAX_FILAS_BUSQUEDA_ENCABEZADO = 30
MAX_FILAS_CABECERA_FICHA = 30

# Campos que deben existir para poder analizar una hoja.
CAMPOS_OBLIGATORIOS = ("documento", "competencia", "juicio", "fecha_juicio")

FORMATOS_FECHA = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y %I:%M:%S %p",
    "%d/%m/%Y",
    "%Y-%m-%d",
)

_FORMATO_FECHA_EXCEL = "yyyy-mm-dd hh:mm"


# ---------------------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------------------

def _normalizar_texto(valor) -> str:
    """Minúsculas, sin tildes ni signos, espacios colapsados."""
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[^a-z0-9 ]+", " ", texto.lower())
    return re.sub(r"\s+", " ", texto).strip()


def _normalizar_documento(valor) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto


def _parsear_fecha(valor) -> datetime | None:
    """Convierte datetime / date / serial de Excel / texto a datetime."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            return from_excel(valor)
        except Exception:
            return None
    texto = str(valor).strip()
    for formato in FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def _es_practica(competencia) -> bool:
    return "etapa practica" in _normalizar_texto(competencia)


def _es_aprobado(juicio) -> bool:
    return _normalizar_texto(juicio) == "aprobado"


# ---------------------------------------------------------------------------
# Detección de encabezado / columnas
# ---------------------------------------------------------------------------

def _clasificar_encabezado(valor) -> str | None:
    """Devuelve el campo al que corresponde un encabezado, o None."""
    texto = _normalizar_texto(valor)
    if not texto:
        return None
    if "funcionario" in texto:
        return "funcionario"
    if texto.startswith("fecha") and "juicio" in texto:
        return "fecha_juicio"
    if texto.startswith("juicio"):
        return "juicio"
    if texto.startswith("resultado") and "aprendizaje" in texto:
        return "resultado"
    if texto.startswith("competencia"):
        return "competencia"
    if texto.startswith("tipo") and ("documento" in texto or "identificacion" in texto):
        return "tipo_documento"
    if any(k in texto for k in ("documento", "identificacion", "cedula")):
        return "documento"
    if texto in ("nombre", "nombres"):
        return "nombre"
    if texto in ("apellido", "apellidos"):
        return "apellidos"
    if texto == "estado":
        return "estado"
    return None


def _detectar_encabezado(ws) -> tuple[int, dict[str, int]] | None:
    """
    Devuelve (fila_encabezado, {campo: columna}) para la fila con más
    campos distintos reconocidos, o None si no hay una hoja utilizable.
    """
    mejor_fila, mejor_cols = 0, {}
    limite = min(ws.max_row, MAX_FILAS_BUSQUEDA_ENCABEZADO)
    for fila in range(1, limite + 1):
        cols: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            campo = _clasificar_encabezado(ws.cell(fila, col).value)
            if campo and campo not in cols:
                cols[campo] = col
        if len(cols) > len(mejor_cols):
            mejor_fila, mejor_cols = fila, cols

    if not all(c in mejor_cols for c in CAMPOS_OBLIGATORIOS):
        return None
    return mejor_fila, mejor_cols


def _detectar_ficha(ws, fila_encabezado: int, nombre_archivo: str) -> tuple[str, str]:
    """
    Busca "Ficha de Caracterización:" en la cabecera y toma el primer valor
    no vacío a su derecha. Devuelve (ficha, metodo).
    """
    limite = min(fila_encabezado - 1, MAX_FILAS_CABECERA_FICHA)
    for fila in range(1, limite + 1):
        for col in range(1, ws.max_column + 1):
            if _normalizar_texto(ws.cell(fila, col).value).startswith("ficha"):
                for col_valor in range(col + 1, ws.max_column + 1):
                    valor = _normalizar_documento(ws.cell(fila, col_valor).value)
                    if valor:
                        return valor, "cabecera"
    return Path(nombre_archivo).stem, "nombre_archivo"


# ---------------------------------------------------------------------------
# Lectura de reportes
# ---------------------------------------------------------------------------

def _leer_reportes(archivos: list[Path]) -> tuple[list[dict], list[dict], list[str]]:
    """
    Devuelve (filas, detalle_columnas, advertencias). Cada fila es un dict
    con los datos de un juicio evaluativo ya normalizados.
    """
    filas: list[dict] = []
    detalle: list[dict] = []
    advertencias: list[str] = []

    for archivo in archivos:
        nombre = Path(archivo).name
        wb = cargar_workbook_compatible(archivo)
        hojas_validas = 0

        for ws in wb.worksheets:
            encabezado = _detectar_encabezado(ws)
            if encabezado is None:
                advertencias.append(
                    f"{nombre} / {ws.title}: no se encontró un encabezado con "
                    f"documento, competencia, juicio y fecha del juicio; hoja omitida."
                )
                continue
            hojas_validas += 1

            fila_enc, cols = encabezado
            ficha, metodo_ficha = _detectar_ficha(ws, fila_enc, nombre)
            if metodo_ficha != "cabecera":
                advertencias.append(
                    f"{nombre} / {ws.title}: no se halló el número de ficha en la "
                    f"cabecera; se usó el nombre del archivo ({ficha})."
                )

            detalle.append({
                "archivo": nombre,
                "hoja": ws.title,
                "ficha": ficha,
                "fila_inicio_datos": fila_enc + 1,
                "columnas": {c: get_column_letter(i) for c, i in cols.items()},
            })

            def valor(fila_vals, campo):
                idx = cols.get(campo)
                return fila_vals[idx - 1] if idx and idx <= len(fila_vals) else None

            sin_fecha_aprobado = 0
            for fila_vals in ws.iter_rows(min_row=fila_enc + 1, values_only=True):
                documento = _normalizar_documento(valor(fila_vals, "documento"))
                if not documento:
                    continue
                juicio = valor(fila_vals, "juicio")
                fecha = _parsear_fecha(valor(fila_vals, "fecha_juicio"))
                aprobado = _es_aprobado(juicio)
                if aprobado and fecha is None:
                    sin_fecha_aprobado += 1

                filas.append({
                    "ficha": ficha,
                    "tipo_documento": valor(fila_vals, "tipo_documento"),
                    "documento": documento,
                    "nombre": valor(fila_vals, "nombre"),
                    "apellidos": valor(fila_vals, "apellidos"),
                    "estado": valor(fila_vals, "estado"),
                    "competencia": valor(fila_vals, "competencia"),
                    "resultado": valor(fila_vals, "resultado"),
                    "juicio": str(juicio).strip() if juicio is not None else "",
                    "aprobado": aprobado,
                    "es_practica": _es_practica(valor(fila_vals, "competencia")),
                    "fecha": fecha,
                    "funcionario": valor(fila_vals, "funcionario"),
                })

            if sin_fecha_aprobado:
                advertencias.append(
                    f"{nombre} / {ws.title}: {sin_fecha_aprobado} juicio(s) APROBADO(S) "
                    f"sin fecha válida; no se pudieron comparar."
                )

        if hojas_validas == 0:
            advertencias.append(f"{nombre}: ninguna hoja del archivo pudo analizarse.")

    return filas, detalle, advertencias


# ---------------------------------------------------------------------------
# Análisis
# ---------------------------------------------------------------------------

def _analizar(filas: list[dict], solo_fecha: bool) -> tuple[list[dict], dict]:
    """
    Agrupa por (ficha, documento), toma la fecha de aprobación de la práctica
    (si hubiera varias filas aprobadas, la más antigua) y marca como
    inconsistencia todo juicio lectivo aprobado con fecha mayor.

    Devuelve (aprendices, estadisticas_por_ficha).
    """
    grupos: dict[tuple[str, str], list[dict]] = {}
    for f in filas:
        grupos.setdefault((f["ficha"], f["documento"]), []).append(f)

    def clave(dt: datetime):
        return dt.date() if solo_fecha else dt

    aprendices: list[dict] = []
    por_ficha: dict[str, dict] = {}

    for (ficha, documento), juicios in grupos.items():
        est = por_ficha.setdefault(ficha, {
            "ficha": ficha,
            "aprendices": 0,
            "con_practica_aprobada": 0,
            "con_inconsistencia": 0,
            "juicios_posteriores": 0,
        })
        est["aprendices"] += 1

        fechas_practica = [
            j["fecha"] for j in juicios
            if j["es_practica"] and j["aprobado"] and j["fecha"]
        ]
        if not fechas_practica:
            continue
        est["con_practica_aprobada"] += 1
        fecha_practica = min(fechas_practica)

        posteriores = [
            j for j in juicios
            if j["aprobado"] and not j["es_practica"] and j["fecha"]
            and clave(j["fecha"]) > clave(fecha_practica)
        ]
        if not posteriores:
            continue

        posteriores.sort(key=lambda j: j["fecha"])
        for j in posteriores:
            j["dias_despues"] = round((j["fecha"] - fecha_practica).total_seconds() / 86400, 1)

        base = posteriores[0]
        aprendices.append({
            "ficha": ficha,
            "tipo_documento": base["tipo_documento"],
            "documento": documento,
            "nombre": base["nombre"],
            "apellidos": base["apellidos"],
            "estado": base["estado"],
            "fecha_practica": fecha_practica,
            "juicios_posteriores": posteriores,
            "primer_posterior": posteriores[0]["fecha"],
            "ultimo_posterior": posteriores[-1]["fecha"],
            "max_dias": max(j["dias_despues"] for j in posteriores),
        })
        est["con_inconsistencia"] += 1
        est["juicios_posteriores"] += len(posteriores)

    aprendices.sort(key=lambda a: (a["ficha"], -len(a["juicios_posteriores"])))
    return aprendices, por_ficha


# ---------------------------------------------------------------------------
# Generación del Excel
# ---------------------------------------------------------------------------

_FUENTE = "Arial"
_RELLENO_ENCABEZADO = PatternFill("solid", start_color="1F4E78")


def _escribir_hoja(ws, encabezados: list[str], filas: list[list], anchos: list[int]):
    ws.append(encabezados)
    for c in range(1, len(encabezados) + 1):
        celda = ws.cell(1, c)
        celda.font = Font(name=_FUENTE, bold=True, color="FFFFFF")
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in filas:
        ws.append(fila)
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name=_FUENTE)
            if isinstance(celda.value, datetime):
                celda.number_format = _FORMATO_FECHA_EXCEL
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions


def _generar_excel(aprendices: list[dict], por_ficha: dict, ruta: Path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen por ficha"
    _escribir_hoja(
        ws,
        ["Ficha", "Aprendices", "Con práctica aprobada", "Con inconsistencia", "Juicios posteriores"],
        [
            [e["ficha"], e["aprendices"], e["con_practica_aprobada"],
             e["con_inconsistencia"], e["juicios_posteriores"]]
            for e in sorted(por_ficha.values(), key=lambda e: e["ficha"])
        ],
        [16, 14, 24, 20, 20],
    )

    ws = wb.create_sheet("Aprendices a revisar")
    _escribir_hoja(
        ws,
        ["Ficha", "Tipo Doc.", "Documento", "Nombre", "Apellidos", "Estado",
         "Fecha aprobación práctica", "# Juicios posteriores",
         "Primer juicio posterior", "Último juicio posterior", "Días máx. después"],
        [
            [a["ficha"], a["tipo_documento"], a["documento"], a["nombre"], a["apellidos"],
             a["estado"], a["fecha_practica"], len(a["juicios_posteriores"]),
             a["primer_posterior"], a["ultimo_posterior"], a["max_dias"]]
            for a in aprendices
        ],
        [12, 10, 15, 24, 26, 18, 24, 14, 22, 22, 16],
    )

    ws = wb.create_sheet("Detalle juicios")
    detalle = []
    for a in aprendices:
        for j in a["juicios_posteriores"]:
            detalle.append([
                a["ficha"], a["documento"], a["nombre"], a["apellidos"], a["estado"],
                j["competencia"], j["resultado"], j["juicio"], j["fecha"],
                a["fecha_practica"], j["dias_despues"], j["funcionario"],
            ])
    _escribir_hoja(
        ws,
        ["Ficha", "Documento", "Nombre", "Apellidos", "Estado", "Competencia",
         "Resultado de Aprendizaje", "Juicio", "Fecha del juicio",
         "Fecha aprobación práctica", "Días después de la práctica", "Funcionario que registró"],
        detalle,
        [12, 15, 22, 24, 16, 50, 50, 12, 20, 22, 14, 40],
    )

    wb.save(ruta)


# ---------------------------------------------------------------------------
# Proceso principal
# ---------------------------------------------------------------------------

def verificar_juicios_posteriores_practica(
    archivos: list[Path],
    salida_dir: Path,
    solo_fecha: bool = False,
) -> dict:
    """
    Cruza los reportes de juicios y genera un Excel con los aprendices que
    tienen juicios lectivos aprobados DESPUÉS de la aprobación de la etapa
    práctica.

    solo_fecha=False: compara fecha y hora (un juicio el mismo día pero
    horas después también se marca). solo_fecha=True: compara solo el día.
    """
    salida_dir.mkdir(parents=True, exist_ok=True)

    filas, detalle_columnas, advertencias = _leer_reportes(archivos)
    aprendices, por_ficha = _analizar(filas, solo_fecha)

    ruta_salida = salida_dir / "Juicios_Posteriores_Practica.xlsx"
    _generar_excel(aprendices, por_ficha, ruta_salida)

    fichas_sin_practica = [
        e["ficha"] for e in por_ficha.values() if e["con_practica_aprobada"] == 0
    ]

    return {
        "archivo_generado": str(ruta_salida),
        "criterio_comparacion": "solo fecha" if solo_fecha else "fecha y hora",
        "total_fichas": len(por_ficha),
        "total_aprendices": sum(e["aprendices"] for e in por_ficha.values()),
        "total_con_practica_aprobada": sum(e["con_practica_aprobada"] for e in por_ficha.values()),
        "total_aprendices_con_inconsistencia": len(aprendices),
        "total_juicios_posteriores": sum(len(a["juicios_posteriores"]) for a in aprendices),
        "resumen_por_ficha": sorted(por_ficha.values(), key=lambda e: e["ficha"]),
        "fichas_sin_practica_aprobada": fichas_sin_practica,
        "aprendices_con_inconsistencia": [
            {
                "ficha": a["ficha"],
                "documento": a["documento"],
                "nombre": f"{a['nombre'] or ''} {a['apellidos'] or ''}".strip(),
                "estado": a["estado"],
                "fecha_practica": a["fecha_practica"].isoformat(sep=" ", timespec="minutes"),
                "juicios_posteriores": len(a["juicios_posteriores"]),
                "ultimo_juicio_posterior": a["ultimo_posterior"].isoformat(sep=" ", timespec="minutes"),
            }
            for a in aprendices[:50]
        ],
        "columnas_detectadas": detalle_columnas,
        "advertencias": advertencias,
    }